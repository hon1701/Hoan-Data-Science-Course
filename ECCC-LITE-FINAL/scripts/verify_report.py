"""Check the saved Word report against the authoritative result summary."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def document_text(doc) -> str:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(document_text(cell))
    return "\n".join(texts)


def main() -> None:
    results_path = ROOT / "outputs/tables/evaluation_summary.json"
    results = json.loads(results_path.read_text(encoding="utf-8"))
    report = ROOT / "reports/BaoCao_NOP.docx"
    provenance = json.loads((ROOT / "reports/report_provenance.json").read_text(encoding="utf-8"))
    assert provenance["evaluation_sha256"] == hashlib.sha256(results_path.read_bytes()).hexdigest()
    assert provenance["report_sha256"] == hashlib.sha256(report.read_bytes()).hexdigest()
    text = document_text(Document(report))
    metrics = results["test"]
    for key in ("average_precision", "roc_auc"):
        assert f"{metrics[key]:.4f}".replace(".", ",") in text, key
    primary = metrics["cost_threshold_metrics"]
    for key in ("precision", "recall", "f1"):
        assert f"{primary[key]:.4f}".replace(".", ",") in text, key
    assert f"{primary['threshold']:.6f}".replace(".", ",") in text
    for row in results["top_p"]:
        assert f"{row['recall_at_k']*100:.2f}%".replace(".", ",") in text
    for marker in ("HƯỚNG DẪN SỬ DỤNG BẢN KHUNG", "TODO", "TBD"):
        assert marker not in text, marker
    print("[OK] Báo cáo khớp hash lần chạy, AP/ROC-AUC, ngưỡng, precision/recall/F1 và Top-p.")


if __name__ == "__main__":
    main()
